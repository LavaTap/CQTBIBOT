"""课表数据模型。"""
from __future__ import annotations

import json
import logging
from dataclasses import dataclass, field
from pathlib import Path

# 保存 JSON 时屏蔽的课程名关键词（不录入磁盘）
_FILTERED_COURSE_KEYWORDS: list[str] = ["AI", "AIGC"]

log = logging.getLogger("schedule.models")


def _should_skip_course(name: str) -> bool:
    """判断课程名是否匹配屏蔽关键词（不保存到 JSON）。"""
    upper = name.upper()
    return any(kw.upper() in upper for kw in _FILTERED_COURSE_KEYWORDS)


class ScheduleError(Exception):
    """课表操作异常。"""
    pass


@dataclass
class Course:
    name: str = ""
    teacher: str = ""
    room: str = ""
    weeks: str = ""
    day: int = 0        # 1=周一 … 7=周日
    period_start: int = 0
    period_end: int = 0

    def week_matches(self, week_num: int) -> bool:
        """判断课程是否在指定周次。weeks 格式如 '2-5,7-9,15'。"""
        if not self.weeks:
            return True
        for part in self.weeks.split(","):
            part = part.strip()
            if "-" in part:
                lo, hi = part.split("-", 1)
                if lo.isdigit() and hi.isdigit() and int(lo) <= week_num <= int(hi):
                    return True
            elif part.isdigit() and int(part) == week_num:
                return True
        return False

    def to_dict(self) -> dict:
        return {
            "name": self.name,
            "teacher": self.teacher,
            "room": self.room,
            "weeks": self.weeks,
            "day": self.day,
            "period_start": self.period_start,
            "period_end": self.period_end,
        }


@dataclass
class Schedule:
    semester: str = ""
    student_id: str = ""
    student_name: str = ""
    fetched_at: str = ""
    courses: list[Course] = field(default_factory=list)

    def _course_key(self, c: Course) -> tuple:
        """课程去重键。"""
        return (
            c.name.strip(), c.teacher.strip(), c.weeks.strip(),
            c.day, c.period_start, c.period_end, c.room.strip(),
        )

    def merge(self, other: Schedule) -> int:
        """将 other 的课程合并进来（按去重键去重），返回新增课程数。"""
        existing = {self._course_key(c) for c in self.courses}
        added = 0
        for c in other.courses:
            if self._course_key(c) not in existing:
                existing.add(self._course_key(c))
                self.courses.append(c)
                added += 1
        return added

    def to_dict(self) -> dict:
        return {
            "semester": self.semester,
            "student_id": self.student_id,
            "student_name": self.student_name,
            "fetched_at": self.fetched_at,
            "courses": [c.to_dict() for c in self.courses],
        }

    def save_json(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        data = self.to_dict()
        # 过滤掉匹配屏蔽关键词的课程（不录入磁盘）
        data["courses"] = [c for c in data["courses"]
                           if not _should_skip_course(c["name"])]
        path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _max_week(self) -> int:
        """从课程周次推算最大周数。"""
        mx = 20
        for c in self.courses:
            for part in c.weeks.split(","):
                part = part.strip()
                if "-" in part:
                    hi = part.split("-", 1)[1]
                    if hi.isdigit():
                        mx = max(mx, int(hi))
                elif part.isdigit():
                    mx = max(mx, int(part))
        return mx

    def save_excel(self, path: Path, week_start: int = 0, week_end: int = 0) -> None:
        """导出按周分 Sheet 的标准课表 Excel。

        每周一个 Sheet，命名「第N周」。Sheet 布局：
            行1:  表头   节次 | 周一 | 周二 | ... | 周日
            行2-6: 5 大节（一大节 1-3 / 二大节 4-5 / 三大节 6-8 / 四大节 9-10 / 五大节 11-12）
            行7:  空行
            行8:  学期: 2025-2026-2

        每个单元格内容（按需拼接，缺则跳过）：
            课程名\n教师\n周次(周)\n教室

        参数 week_start/week_end=0 表示导出全部周（按课程实际周次自动推算上限）。
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)

        header_font = Font(bold=True, size=11)
        header_fill = PatternFill("solid", fgColor="D6E4F0")
        period_font = Font(bold=True, size=10)
        period_fill = PatternFill("solid", fgColor="EAF1F8")
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)

        period_rows = [
            (1, 3, "一大节\n1-3节"),
            (4, 5, "二大节\n4-5节"),
            (6, 8, "三大节\n6-8节"),
            (9, 10, "四大节\n9-10节"),
            (11, 12, "五大节\n11-12节"),
        ]
        day_names = ["节次", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]

        max_w = self._max_week()
        ws_start = week_start if week_start > 0 else 1
        ws_end = week_end if week_end > 0 else max_w
        if ws_start > ws_end:
            ws_start, ws_end = ws_end, ws_start

        def _cell_text(course: Course) -> str:
            parts = [course.name]
            if course.teacher:
                parts.append(course.teacher)
            if course.weeks:
                parts.append(f"{course.weeks}(周)")
            if course.room:
                parts.append(course.room)
            return "\n".join(parts)

        for w in range(ws_start, ws_end + 1):
            ws = wb.create_sheet(title=f"第{w}周")

            for ci, name in enumerate(day_names, 1):
                cell = ws.cell(row=1, column=ci, value=name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = border

            for ri, (ps, pe, label) in enumerate(period_rows, start=2):
                cell = ws.cell(row=ri, column=1, value=label)
                cell.font = period_font
                cell.fill = period_fill
                cell.alignment = center
                cell.border = border
                for day in range(1, 8):
                    col = day + 1
                    matches = [
                        c for c in self.courses
                        if c.day == day
                        and c.period_start <= pe and c.period_end >= ps
                        and c.week_matches(w)
                    ]
                    cell = ws.cell(row=ri, column=col)
                    cell.border = border
                    cell.alignment = left_top
                    if matches:
                        cell.value = "\n---\n".join(_cell_text(m) for m in matches)
                        ws.row_dimensions[ri].height = max(
                            ws.row_dimensions[ri].height or 60, 60,
                        )

            ws.cell(row=8, column=1, value="学期:").font = Font(bold=True)
            ws.cell(row=8, column=2, value=self.semester)

            ws.column_dimensions["A"].width = 10
            for ci in range(2, 9):
                ws.column_dimensions[chr(64 + ci)].width = 22

        wb.save(path)
